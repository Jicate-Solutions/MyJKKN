"""
HR Recruitment — CVViZ job postings importer.

Reads the CVViZ "All Jobs" export (cvviz_all_jobs.xlsx, 55 columns, one row per
CVViZ job) and loads it into public.hr_recruitment_jobs.

Scope decisions (Director, 2026-10-01):
- Skip: jobs deleted in CVViZ, US-located jobs, Jicate Solutions jobs, and the
  two CVViZ jobs that duplicate live MyJKKN jobs (83946, 81212).
- Every imported job lands as status='draft', is_public=false — HR reopens the
  real vacancies one by one; nothing reaches /careers on import.
- Arts & Science college (Aided vs Self) is decided by department: a department
  that exists in only one of the two colleges picks that college; one that
  exists in both, or no department, defaults to Self (flagged for review).

Requires migration 20261001090128_hr_recruitment_jobs_external_source.sql
(external_source / external_id / external_url / external_meta) before --apply.
Re-running is safe: rows already imported (same external_source + external_id)
are skipped.

Usage:
    python scripts/import-cvviz-jobs.py <xlsx> --preview <out.xlsx>
        → dry-run: resolves departments read-only, writes the review sheet
    python scripts/import-cvviz-jobs.py <xlsx> --apply
        → inserts into production via Supabase REST (service role)

Reads NEXT_PUBLIC_SUPABASE_URL and SUPABASE_SERVICE_ROLE_KEY from .env.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import urllib.error
import urllib.parse
import urllib.request
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl

REPO_ROOT = Path(__file__).resolve().parents[1]
ENV_PATH = REPO_ROOT / ".env"
SOURCE = "cvviz"

ARTS_AIDED = "a33138b6-4eea-4675-941f-1071bf88b127"
ARTS_SELF = "b0b8a724-7c65-4f07-8047-2a38e8100ad5"
JICATE = "479eac7f-3e5b-479e-bd91-dee9e0186b9b"
DUPLICATES_OF_LIVE_JOBS = {"83946", "81212"}

# CVViZ department prefix → MyJKKN department_name, for Arts & Science rows.
ARTS_DEPT_ALIASES = {
    "Maths": "Mathematics",
    "Textile and Fashion Designing": "Textile Fashion Designing",
}

EDUCATION_LEVEL = {
    "High school": "high_school",
    "Bachelor": "bachelors",
    "Master": "masters",
    "Doctorate / PhD": "phd",
}

PROFICIENCY = {"basic": "basic", "intermediate": "intermediate", "pro": "pro", "expert": "expert"}

LEADERSHIP_RE = re.compile(r"\b(principal|chief|ceo|coo|director)\b", re.I)
TEACHING_RE = re.compile(
    r"prof|lecturer|reader|tutor|teacher|educator|trainer|coach|"
    r"\bb ?t assistant\b|graduate assistant|facilitator|faculty",
    re.I,
)


# --- Env / REST -------------------------------------------------------------
def load_env() -> dict[str, str]:
    env: dict[str, str] = {}
    with open(ENV_PATH) as f:
        for line in f:
            line = line.rstrip("\n\r")
            if not line or line.startswith("#") or "=" not in line:
                continue
            k, v = line.split("=", 1)
            env[k.strip()] = v.strip().strip('"').strip("'")
    return env


class Supabase:
    def __init__(self, url: str, key: str):
        self.url = url.rstrip("/")
        self.headers = {
            "apikey": key,
            "Authorization": f"Bearer {key}",
            "Content-Type": "application/json",
        }

    def get(self, path: str) -> list[dict[str, Any]]:
        req = urllib.request.Request(f"{self.url}/rest/v1/{path}", headers=self.headers)
        with urllib.request.urlopen(req) as r:
            return json.loads(r.read() or "[]")

    def insert_ignore_duplicates(self, table: str, rows: list[dict[str, Any]], on_conflict: str) -> int:
        req = urllib.request.Request(
            f"{self.url}/rest/v1/{table}?on_conflict={on_conflict}",
            data=json.dumps(rows).encode(),
            method="POST",
            headers={**self.headers, "Prefer": "resolution=ignore-duplicates,return=representation"},
        )
        try:
            with urllib.request.urlopen(req) as r:
                return len(json.loads(r.read() or "[]"))
        except urllib.error.HTTPError as e:
            raise RuntimeError(f"{e.code} {e.read().decode()[:500]}") from None


# --- Parsing helpers --------------------------------------------------------
def clean(v: Any) -> str | None:
    if v is None:
        return None
    s = re.sub(r"\s+", " ", str(v)).strip()
    return s or None


def to_int(v: Any) -> int | None:
    try:
        return int(float(v)) if v not in (None, "") else None
    except (TypeError, ValueError):
        return None


def parse_ts(v: Any) -> str | None:
    """CVViZ dates look like 'Tue Jul 21 2026 09:46:04' (UTC)."""
    s = clean(v)
    if not s:
        return None
    dt = datetime.strptime(s, "%a %b %d %Y %H:%M:%S").replace(tzinfo=timezone.utc)
    return dt.isoformat()


def split_list(v: Any) -> list[str]:
    return [p.strip() for p in str(v or "").split(";") if p.strip()]


def parse_skills(v: Any) -> list[dict[str, str]]:
    skills, seen = [], set()
    for part in split_list(v):
        bits = [b.strip() for b in part.split("|")]
        name = clean(bits[0])
        if not name or name.lower() in seen:
            continue
        seen.add(name.lower())
        level = PROFICIENCY.get((bits[1] if len(bits) > 1 else "").lower(), "intermediate")
        kind = "nice_to_have" if len(bits) > 2 and bits[2].lower().startswith("nice") else "required"
        skills.append({"name": name, "type": kind, "proficiency": level})
    return skills


def role_category(title: str, job_function: str | None) -> str:
    text = f"{title} {job_function or ''}"
    if LEADERSHIP_RE.search(title):
        return "senior_leadership"
    if TEACHING_RE.search(text) and not re.search(r"lab assis|librar", title, re.I):
        return "teaching_faculty"
    return "non_teaching"


def experience_text(lo: int | None, hi: int | None) -> str | None:
    if lo is None and hi is None:
        return None
    if hi is None:
        return f"{lo}+ years"
    if lo is None:
        return f"Up to {hi} years"
    return f"{lo}–{hi} years"


# --- Mapping ----------------------------------------------------------------
def skip_reason(r: dict[str, Any]) -> str | None:
    if r["Deleted in CVViZ"] == "Yes":
        return "deleted in CVViZ"
    if r["Country"] != "India":
        return "US-located"
    if r["MyJKKN institution_id"] == JICATE:
        return "Jicate Solutions"
    if str(r["CVViZ Job ID"]) in DUPLICATES_OF_LIVE_JOBS:
        return "duplicate of live MyJKKN job"
    return None


def resolve_institution_and_dept(
    r: dict[str, Any], depts: dict[str, dict[str, str]]
) -> tuple[str, str | None, list[str]]:
    """Returns (institution_id, department_name or None, flags)."""
    flags: list[str] = []
    inst = r["MyJKKN institution_id"]
    suggested = clean(r["Suggested MyJKKN Department"])
    dept = None
    if suggested and suggested != "(institution level / no dept)":
        dept = re.sub(r"\s*\[.*\]$", "", suggested)

    if not inst:
        # Arts & Science rows CVViZ left undecided (Aided vs Self), plus rows
        # whose CVViZ department was deleted.
        if not dept:
            prefix = clean((r["CVViZ Department"] or "").split("-")[0])
            dept = ARTS_DEPT_ALIASES.get(prefix, prefix) if prefix else None
        in_aided = bool(dept) and dept.lower() in depts[ARTS_AIDED]
        in_self = bool(dept) and dept.lower() in depts[ARTS_SELF]
        if in_aided and not in_self:
            inst = ARTS_AIDED
        else:
            inst = ARTS_SELF
            if not in_self:
                flags.append("college defaulted to A&S Self (no department)")
                dept = None
            elif in_aided:
                flags.append("college defaulted to A&S Self (department in both)")

    if r["Mapping Note"] and "Sresakthimayeil" in r["Mapping Note"]:
        flags.append("Sresakthimayeil job mapped to Nursing (CNR)")
    if dept and dept.lower() not in depts.get(inst, {}):
        flags.append(f"department '{dept}' not found in college; left blank")
        dept = None
    return inst, dept, flags


def map_row(r: dict[str, Any], depts: dict[str, dict[str, str]]) -> tuple[dict[str, Any], list[str]]:
    inst, dept_name, flags = resolve_institution_and_dept(r, depts)
    title = clean(r["Title"]) or "Untitled"
    category = role_category(title, clean(r["Job Function"]))

    lo, hi = to_int(r["Min Exp (yrs)"]), to_int(r["Max Exp (yrs)"])
    if lo is not None and hi is not None and lo > hi:
        hi = None

    # CVViZ's currency/interval defaults are unreliable (276 Tamil Nadu jobs
    # say "USD / Yearly" with 20000–30000): every imported job is in India, so
    # amounts are read as INR per month. The as-entered values go to
    # external_meta.salary_as_entered.
    smin, smax = to_int(r["Salary Min"]), to_int(r["Salary Max"])
    if smin is not None and smax is not None and smin > smax:
        smin, smax = smax, smin
    if smin is not None or smax is not None:
        if (smin or smax or 0) < 1000:
            flags.append("salary looks hourly in CVViZ; left blank")
            smin = smax = None

    quals = clean(r["Qualifications"])
    requirements: dict[str, Any] = {}
    if quals and quals.upper() != "ANY":
        requirements["qualifications"] = [quals]
    skills = parse_skills(r["Skills (name | level | required)"])
    if skills:
        requirements["skills"] = skills
    exp = experience_text(lo, hi)
    if exp:
        requirements["experience"] = exp

    created = parse_ts(r["Created At"])
    meta = {
        "job_code": clean(r["Job Code"]),
        "status": clean(r["Status"]),
        "approval_status": clean(r["Approval Status"]),
        "cvviz_department": clean(r["CVViZ Department"]),
        "mapping_note": clean(r["Mapping Note"]),
        "job_function": clean(r["Job Function"]),
        "employer_type": clean(r["Employer Type"]),
        "salary_as_entered": {
            "min": r["Salary Min"], "max": r["Salary Max"],
            "currency": clean(r["Currency"]), "interval": clean(r["Pay Interval"]),
        },
        "remote": r["Remote"] == "Yes",
        "job_validity_days": to_int(r["Job Validity (days)"]),
        "screening_questions": split_list(r["Screening Questions"]),
        "feedback_type": clean(r["Feedback Type"]),
        "feedback_criteria": split_list(r["Feedback Criteria"]),
        "benchmark_resumes": split_list(r["Benchmark Resumes"]),
        "assigned_recruiters": split_list(r["Assigned Recruiters"]),
        "hiring_manager": clean(r["Hiring Manager"]),
        "approvers": clean(r["Approvers (responded)"]),
        "candidates_total": to_int(r["Candidates Total"]),
        "stage_counts": clean(r["Stage Counts"]),
        "job_notes": clean(r["Job Notes"]),
        "published_on": [p.strip() for p in str(r["Published On"] or "").split(",") if p.strip()],
        "created_by": clean(r["Created By"]),
        "updated_by": clean(r["Updated By"]),
        "last_evaluated": parse_ts(r["Last Evaluated"]),
        "closed_at": parse_ts(r["Closed At"]),
        "close_reason": clean(r["Close Reason"]),
    }

    row = {
        "institution_id": inst,
        "department_id": depts.get(inst, {}).get(dept_name.lower()) if dept_name else None,
        "title": title,
        "role_category": category,
        "description": clean(r["Description (text)"]),
        "requirements": requirements,
        "min_monthly_salary": smin,
        "max_monthly_salary": smax,
        "salary_currency": "INR",
        "salary_duration": "per_month",
        "display_salary": r["Show Salary on Career Page"] == "Yes" and (smin is not None or smax is not None),
        "status": "draft",
        "is_public": False,
        "posted_at": created,
        "job_code": clean(r["Job Code"]),
        "job_type": "full_time",
        "industry": clean(r["Industry"]),
        "employer_type": "educational",
        "country": "India",
        "state": "Tamil Nadu",
        "city": "Kumarapalayam",
        "zip_code": clean(r["ZIP"]),
        "education_level": EDUCATION_LEVEL.get(clean(r["Education Level"]) or ""),
        "min_experience_years": lo,
        "max_experience_years": hi,
        "created_at": created,
        "updated_at": parse_ts(r["Updated At"]) or created,
        "external_source": SOURCE,
        "external_id": str(r["CVViZ Job ID"]),
        "external_url": clean(r["Share URL"]),
        "external_meta": meta,
    }
    return row, flags


# --- Main -------------------------------------------------------------------
def load_rows(xlsx: Path) -> list[dict[str, Any]]:
    ws = openpyxl.load_workbook(xlsx, read_only=True).worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = list(next(it))
    return [dict(zip(header, r)) for r in it if any(v is not None for v in r)]


def load_departments(sb: Supabase) -> dict[str, dict[str, str]]:
    out: dict[str, dict[str, str]] = {}
    for d in sb.get("departments?select=id,department_name,institution_id&limit=5000"):
        out.setdefault(d["institution_id"], {})[d["department_name"].strip().lower()] = d["id"]
    return out


def write_preview(path: Path, kept: list[tuple[dict, list[str]]], skipped: list[tuple[dict, str]],
                  inst_names: dict[str, str], dept_names: dict[str, str]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "To import"
    cols = ["CVViZ Job ID", "Title", "College", "Department", "Job category", "Status", "Public",
            "Salary min (INR/month)", "Salary max", "Experience", "Education level",
            "Qualifications", "Skills", "Posted at", "Job code", "Review flags"]
    ws.append(cols)
    for row, flags in kept:
        req = row["requirements"]
        ws.append([
            row["external_id"], row["title"], inst_names.get(row["institution_id"], row["institution_id"]),
            dept_names.get(row["department_id"] or "", ""), row["role_category"], row["status"],
            "Yes" if row["is_public"] else "No", row["min_monthly_salary"], row["max_monthly_salary"],
            req.get("experience"), row["education_level"], "; ".join(req.get("qualifications", [])),
            "; ".join(f"{s['name']} ({s['proficiency']}, {s['type']})" for s in req.get("skills", [])),
            (row["posted_at"] or "")[:10], row["job_code"], "; ".join(flags),
        ])
    ws2 = wb.create_sheet("Skipped")
    ws2.append(["CVViZ Job ID", "Title", "CVViZ status", "Reason"])
    for r, reason in skipped:
        ws2.append([r["CVViZ Job ID"], clean(r["Title"]), r["Status"], reason])
    for sheet in (ws, ws2):
        sheet.freeze_panes = "A2"
        for c in sheet[1]:
            c.font = openpyxl.styles.Font(bold=True)
    wb.save(path)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", type=Path)
    ap.add_argument("--preview", type=Path, help="write the review workbook here (dry-run)")
    ap.add_argument("--apply", action="store_true", help="insert into production")
    args = ap.parse_args()

    env = load_env()
    sb = Supabase(env["NEXT_PUBLIC_SUPABASE_URL"], env["SUPABASE_SERVICE_ROLE_KEY"])
    depts = load_departments(sb)

    kept, skipped = [], []
    for r in load_rows(args.xlsx):
        reason = skip_reason(r)
        if reason:
            skipped.append((r, reason))
        else:
            kept.append(map_row(r, depts))

    codes = [row["job_code"] for row, _ in kept if row["job_code"]]
    dupes = {c for c in codes if codes.count(c) > 1}
    for row, flags in kept:
        if row["job_code"] in dupes:
            row["job_code"] = f"{row['job_code']}-{row['external_id']}"
            flags.append("job code made unique")

    print(f"{len(kept)} to import, {len(skipped)} skipped")

    if args.preview:
        insts = {i["id"]: i["name"] for i in sb.get("institutions?select=id,name")}
        dnames = {d["id"]: d["department_name"] for d in sb.get("departments?select=id,department_name&limit=5000")}
        write_preview(args.preview, kept, skipped, insts, dnames)
        print(f"preview written to {args.preview}")

    if args.apply:
        rows = [row for row, _ in kept]
        inserted = 0
        for i in range(0, len(rows), 50):
            inserted += sb.insert_ignore_duplicates(
                "hr_recruitment_jobs", rows[i:i + 50], "external_source,external_id"
            )
        print(f"inserted {inserted} (already present: {len(rows) - inserted})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
